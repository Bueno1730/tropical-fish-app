# -*- coding: utf-8 -*-
"""Genera el documento Word de presentacion de la app Tropical Fish."""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = os.path.dirname(os.path.abspath(__file__))
REPO = "https://github.com/Bueno1730/tropical-fish-app"
APK = REPO + "/releases/download/v1.0/tropical_fish.apk"
RELEASES = REPO + "/releases"
WEB = "https://tropical-fish.onrender.com"


def hipervinculo(parrafo, texto, url):
    """Agrega un hipervinculo cliqueable a un parrafo."""
    parte = parrafo.part
    r_id = parte.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single')
    rPr.append(color); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = texto
    run.append(t)
    link.append(run)
    parrafo._p.append(link)
    return link


doc = Document()

# Titulo
t = doc.add_heading('Aplicacion Movil "Tropical Fish"', 0)
sub = doc.add_paragraph('Conversion del sistema web a aplicacion Android')
sub.runs[0].font.size = Pt(13)
sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# 1. Introduccion
doc.add_heading('1. Introduccion', 1)
p = doc.add_paragraph(
    'El sistema de gestion del Acuario Tropical Fish es una aplicacion web '
    '(PHP/MVC) publicada en Internet. Para utilizarla como aplicacion instalable '
    'en telefonos Android se desarrollo un contenedor nativo (WebView) que '
    'empaqueta el sistema web dentro de un archivo APK.')
p2 = doc.add_paragraph('Sistema web en produccion: ')
hipervinculo(p2, WEB, WEB)

# 2. Como se convierte
doc.add_heading('2. Como se convierte la web en aplicacion', 1)
doc.add_paragraph(
    'La conversion no reescribe el sistema: la app es una "carcasa" nativa '
    'que muestra la web dentro del telefono. Sus componentes son:', style=None)
items = [
    'Activity principal (MainActivity.java): crea un WebView de Android a '
    'pantalla completa y carga la URL del sistema al iniciar la app.',
    'WebView con JavaScript, almacenamiento DOM y cookies habilitados, por '
    'lo que el sistema funciona igual que en un navegador.',
    'Funciones nativas integradas: selector de archivos del telefono para la '
    'importacion CSV, gestor de descargas para los reportes exportados, boton '
    '"atras" que navega dentro del historial web y menu para recargar o '
    'cambiar el servidor.',
    'Manifiesto Android (AndroidManifest.xml): declara el icono, el nombre, '
    'los permisos de red y el tema visual de la app.',
]
for it in items:
    doc.add_paragraph(it, style='List Bullet')

doc.add_paragraph(
    'El APK se compilo sin Gradle ni Android Studio, usando directamente las '
    'build-tools del Android SDK mediante el script build_apk.py:')
pasos = [
    'aapt2 compile / link  ->  compila los recursos y el manifiesto',
    'javac                 ->  compila el codigo Java',
    'd8                    ->  convierte el bytecode a formato DEX (Android)',
    'zipalign              ->  alinea el paquete para ejecucion eficiente',
    'apksigner             ->  firma el APK para que Android lo instale',
]
for s in pasos:
    doc.add_paragraph(s, style='List Bullet')

# 3. Captura
doc.add_heading('3. Captura de pantalla de la app funcionando', 1)
doc.add_paragraph(
    'La aplicacion instalada en el telefono abre directamente el login del '
    'sistema, servido desde la web en la nube:')
img_p = doc.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.add_run().add_picture(os.path.join(RAIZ, 'captura_app.png'), width=Inches(2.9))

# 4. Instalacion
doc.add_heading('4. Instalacion', 1)
doc.add_paragraph('Codigo fuente y APK en el repositorio de GitHub:')
pr = doc.add_paragraph('Repositorio: ')
hipervinculo(pr, REPO, REPO)
pa = doc.add_paragraph('Descarga directa del APK: ')
hipervinculo(pa, APK, APK)

doc.add_paragraph('Pasos para instalar en el telefono:')
for s in [
    'Descargar tropical_fish.apk desde el enlace anterior.',
    'Abrir el archivo en el telefono; Android pedira permitir "instalar '
    'apps de origenes desconocidos" la primera vez.',
    'Abrir "Tropical Fish": la app conecta automaticamente con el sistema.',
]:
    doc.add_paragraph(s, style='List Number')
doc.add_paragraph('Tambien se puede instalar por USB con: adb install -r tropical_fish.apk')

salida = os.path.join(RAIZ, 'Tropical_Fish_App.docx')
doc.save(salida)
print('Generado:', salida)
